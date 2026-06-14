"""Workbook-to-MongoDB seed loader for reset/upsert modes (T-015)."""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook

from app.core.config import Settings
from app.seed.mongo import SeedMongoError, reset_demo_collections, seed_mongo_session, upsert_documents
from app.seed.template_spec import SheetSpec, demo_seed_collections, sheets_in_load_order
from app.seed.validator import ValidationResult, validate_workbook
from app.users.security import hash_password

SeedMode = Literal["reset", "upsert"]

BOOLEAN_COLUMNS: frozenset[str] = frozenset(
    {
        "is_demo_user",
        "is_active",
        "is_duplicate_candidate",
        "is_eligible",
        "lead_created",
        "limit_exceeded",
        "callback_available",
        "freeze_simulated",
        "expected_escalation",
        "active",
    }
)

TIMESTAMP_COLUMNS: frozenset[str] = frozenset({"created_at", "updated_at"})


class SeedLoadError(Exception):
    """Seed loader failure (validation, env guard, or Mongo)."""


@dataclass
class CollectionLoadStats:
    """Per-collection seed operation counts."""

    deleted: int = 0
    inserted: int = 0
    updated: int = 0
    skipped: int = 0


@dataclass
class SeedLoadResult:
    """Aggregate seed load outcome."""

    mode: SeedMode
    ok: bool
    validation: ValidationResult
    collections: dict[str, CollectionLoadStats] = field(default_factory=dict)
    database_name: str = ""
    env_label: str = "local"
    error_message: str | None = None

    @property
    def total_deleted(self) -> int:
        return sum(stats.deleted for stats in self.collections.values())

    @property
    def total_inserted(self) -> int:
        return sum(stats.inserted for stats in self.collections.values())

    @property
    def total_updated(self) -> int:
        return sum(stats.updated for stats in self.collections.values())


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_empty(value: Any) -> bool:
    return _cell_str(value) == ""


def _parse_header_row(row: tuple[Any, ...]) -> list[str]:
    return [_cell_str(cell) for cell in row if not _is_empty(cell)]


def _row_dict(headers: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    data: dict[str, Any] = {}
    for index, header in enumerate(headers):
        if index < len(row):
            data[header] = row[index]
        else:
            data[header] = None
    return data


def _parse_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    normalized = _cell_str(value).lower()
    if normalized in {"true", "1", "yes", "y"}:
        return True
    if normalized in {"false", "0", "no", "n"}:
        return False
    raise ValueError(f"invalid boolean value: {value!r}")


def _parse_timestamp(value: Any) -> datetime:
    if isinstance(value, datetime):
        if value.tzinfo is None:
            return value.replace(tzinfo=UTC)
        return value.astimezone(UTC)
    text = _cell_str(value)
    if text.endswith("Z"):
        text = text[:-1] + "+00:00"
    parsed = datetime.fromisoformat(text)
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=UTC)
    return parsed.astimezone(UTC)


def _coerce_value(column: str, value: Any) -> Any:
    if _is_empty(value):
        return None
    if column in BOOLEAN_COLUMNS:
        return _parse_bool(value)
    if column in TIMESTAMP_COLUMNS:
        return _parse_timestamp(value)
    if column in {"amount", "emi_amount", "outstanding_balance", "max_amount", "estimated_tokens"}:
        if isinstance(value, (int, float)):
            return value
        return float(_cell_str(value))
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def row_to_document(spec: SheetSpec, row: dict[str, Any]) -> dict[str, Any]:
    """Transform one Excel row dict into a MongoDB document."""
    document: dict[str, Any] = {}
    all_columns = spec.required_columns + spec.optional_columns

    for column in all_columns:
        raw = row.get(column)
        if _is_empty(raw):
            continue
        if column in spec.json_columns:
            document[column] = json.loads(_cell_str(raw))
            continue
        document[column] = _coerce_value(column, raw)

    if spec.sheet_name == "users":
        plain_password = document.pop("password", None)
        if plain_password is None:
            raise SeedLoadError("users row missing password after validation")
        document["password_hash"] = hash_password(str(plain_password))

    now = datetime.now(UTC)
    if "created_at" not in document:
        document["created_at"] = now
    if "updated_at" not in document:
        document["updated_at"] = now

    return document


def read_workbook_rows(workbook_path: Path) -> dict[str, list[dict[str, Any]]]:
    """Read non-empty data rows from all sheets in the workbook."""
    from app.seed.template_spec import SHEET_SPECS_BY_NAME

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    rows_by_sheet: dict[str, list[dict[str, Any]]] = {}
    try:
        for sheet_name in SHEET_SPECS_BY_NAME:
            if sheet_name not in workbook.sheetnames:
                rows_by_sheet[sheet_name] = []
                continue
            worksheet = workbook[sheet_name]
            sheet_rows = list(worksheet.iter_rows(values_only=True))
            if len(sheet_rows) <= 1:
                rows_by_sheet[sheet_name] = []
                continue
            headers = _parse_header_row(sheet_rows[0])
            data_rows: list[dict[str, Any]] = []
            for row in sheet_rows[1:]:
                if all(_is_empty(cell) for cell in row):
                    continue
                data_rows.append(_row_dict(headers, row))
            rows_by_sheet[sheet_name] = data_rows
    finally:
        workbook.close()
    return rows_by_sheet


def _init_collection_stats() -> dict[str, CollectionLoadStats]:
    return {name: CollectionLoadStats() for name in demo_seed_collections()}


def run_seed(
    mode: SeedMode,
    workbook_path: Path,
    settings: Settings,
    *,
    env_label: str = "local",
    force_reset: bool = False,
) -> SeedLoadResult:
    """Validate workbook then reset or upsert demo collections."""
    validation = validate_workbook(workbook_path)
    if not validation.ok:
        return SeedLoadResult(
            mode=mode,
            ok=False,
            validation=validation,
            error_message="workbook validation failed",
        )

    if mode == "reset" and env_label != "local" and not force_reset:
        return SeedLoadResult(
            mode=mode,
            ok=False,
            validation=validation,
            env_label=env_label,
            error_message=(
                f"reset on env={env_label} requires --force; use upsert for non-destructive reload"
            ),
        )

    rows_by_sheet = read_workbook_rows(workbook_path)
    collection_stats = _init_collection_stats()

    try:
        with seed_mongo_session(settings) as (_client, database):
            if mode == "reset":
                deleted_counts = reset_demo_collections(database)
                for collection_name, count in deleted_counts.items():
                    collection_stats[collection_name].deleted = count

            for spec in sheets_in_load_order():
                sheet_rows = rows_by_sheet.get(spec.sheet_name, [])
                if not sheet_rows:
                    continue

                documents: list[dict[str, Any]] = []
                skipped = 0
                for row in sheet_rows:
                    try:
                        documents.append(row_to_document(spec, row))
                    except (ValueError, SeedLoadError):
                        skipped += 1

                stats = collection_stats[spec.collection_target]
                stats.skipped += skipped
                if not documents:
                    continue

                inserted, updated = upsert_documents(
                    database[spec.collection_target],
                    documents,
                    spec.stable_id_column,
                )
                stats.inserted += inserted
                stats.updated += updated

            return SeedLoadResult(
                mode=mode,
                ok=True,
                validation=validation,
                collections=collection_stats,
                database_name=settings.mongodb_db_name,
                env_label=env_label,
            )
    except SeedMongoError as exc:
        return SeedLoadResult(
            mode=mode,
            ok=False,
            validation=validation,
            collections=collection_stats,
            env_label=env_label,
            error_message=str(exc),
        )


def format_seed_summary(result: SeedLoadResult) -> str:
    """Render dry per-collection seed summary (no secrets)."""
    lines = [
        f"Seed {result.mode} complete "
        f"(env={result.env_label}, db={result.database_name})"
    ]
    for collection_name in demo_seed_collections():
        stats = result.collections.get(collection_name, CollectionLoadStats())
        if result.mode == "reset":
            lines.append(
                f"  {collection_name}: deleted={stats.deleted} "
                f"inserted={stats.inserted} updated={stats.updated}"
            )
        else:
            lines.append(
                f"  {collection_name}: inserted={stats.inserted} updated={stats.updated}"
            )
    if result.mode == "reset":
        lines.append(
            f"Total: deleted={result.total_deleted} "
            f"inserted={result.total_inserted} updated={result.total_updated}"
        )
    else:
        lines.append(
            f"Total: inserted={result.total_inserted} updated={result.total_updated}"
        )
    return "\n".join(lines)
