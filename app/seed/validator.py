"""Read-only validation for Samadhan demo seed Excel workbooks (T-014)."""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from app.seed.template_spec import (
    SHEET_SPECS_BY_NAME,
    SheetSpec,
    header_columns,
    sheets_required_for_validate,
)

ValidationLevel = Literal["error", "warning"]


@dataclass(frozen=True)
class ValidationIssue:
    """Single validation finding."""

    level: ValidationLevel
    sheet: str | None
    row: int | None
    column: str | None
    code: str
    message: str

    def format_line(self) -> str:
        parts = [f"[{self.level.upper()}]", f"code={self.code}"]
        if self.sheet:
            parts.append(f"sheet={self.sheet}")
        if self.row is not None:
            parts.append(f"row={self.row}")
        if self.column:
            parts.append(f"column={self.column}")
        parts.append(self.message)
        return " ".join(parts)


@dataclass
class ValidationResult:
    """Aggregate validation outcome."""

    ok: bool
    issues: list[ValidationIssue] = field(default_factory=list)
    workbook_path: Path = field(default_factory=Path)
    sheets_checked: int = 0
    rows_checked: int = 0

    @property
    def errors(self) -> list[ValidationIssue]:
        return [issue for issue in self.issues if issue.level == "error"]


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_empty(value: Any) -> bool:
    return _cell_str(value) == ""


def _parse_header_row(row: tuple[Any, ...]) -> list[str]:
    return [_cell_str(cell) for cell in row if not _is_empty(cell)]


def _row_dict(headers: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    data: dict[str, Any] = {}
    for index, header in enumerate(headers):
        if index < len(row):
            data[header] = row[index]
        else:
            data[header] = None
    return data


def _validate_headers(
    spec: SheetSpec,
    headers: list[str],
    issues: list[ValidationIssue],
) -> None:
    header_set = set(headers)
    for column in spec.required_columns:
        if column not in header_set:
            issues.append(
                ValidationIssue(
                    level="error",
                    sheet=spec.sheet_name,
                    row=1,
                    column=column,
                    code="MISSING_COLUMN",
                    message=f"required column '{column}' is missing from header row",
                )
            )


def _validate_row(
    spec: SheetSpec,
    headers: list[str],
    row_number: int,
    row: tuple[Any, ...],
    issues: list[ValidationIssue],
) -> None:
    data = _row_dict(headers, row)

    stable_id = spec.stable_id_column
    stable_value = data.get(stable_id)
    if _is_empty(stable_value):
        issues.append(
            ValidationIssue(
                level="error",
                sheet=spec.sheet_name,
                row=row_number,
                column=stable_id,
                code="EMPTY_STABLE_ID",
                message=f"stable id column '{stable_id}' must not be empty",
            )
        )

    for column in spec.required_columns:
        if _is_empty(data.get(column)):
            if column == "password" and spec.sheet_name == "users":
                message = "password must not be empty"
            else:
                message = f"required field '{column}' must not be empty"
            issues.append(
                ValidationIssue(
                    level="error",
                    sheet=spec.sheet_name,
                    row=row_number,
                    column=column,
                    code="EMPTY_REQUIRED_FIELD",
                    message=message,
                )
            )

    for column, allowed in spec.allowed_enums.items():
        raw = data.get(column)
        if _is_empty(raw):
            continue
        value = _cell_str(raw)
        if value not in allowed:
            allowed_text = ", ".join(allowed)
            issues.append(
                ValidationIssue(
                    level="error",
                    sheet=spec.sheet_name,
                    row=row_number,
                    column=column,
                    code="INVALID_ENUM",
                    message=f"invalid enum value (allowed: {allowed_text})",
                )
            )

    for column in spec.json_columns:
        raw = data.get(column)
        if _is_empty(raw):
            continue
        try:
            json.loads(_cell_str(raw))
        except json.JSONDecodeError:
            issues.append(
                ValidationIssue(
                    level="error",
                    sheet=spec.sheet_name,
                    row=row_number,
                    column=column,
                    code="INVALID_JSON",
                    message="value must be valid JSON",
                )
            )

    if spec.sheet_name == "users" and _cell_str(data.get("role")) == "customer":
        if _is_empty(data.get("customer_id")):
            issues.append(
                ValidationIssue(
                    level="error",
                    sheet=spec.sheet_name,
                    row=row_number,
                    column="customer_id",
                    code="MISSING_CUSTOMER_LINK",
                    message="customer role requires customer_id",
                )
            )


def _validate_duplicate_ids(
    spec: SheetSpec,
    id_values: list[tuple[int, str]],
    issues: list[ValidationIssue],
) -> None:
    seen: dict[str, int] = {}
    for row_number, value in id_values:
        if value in seen:
            issues.append(
                ValidationIssue(
                    level="error",
                    sheet=spec.sheet_name,
                    row=row_number,
                    column=spec.stable_id_column,
                    code="DUPLICATE_ID",
                    message=(
                        f"duplicate {spec.stable_id_column} '{value}' "
                        f"(first seen on row {seen[value]})"
                    ),
                )
            )
        else:
            seen[value] = row_number


def _resolve_app_root(workbook_path: Path) -> Path:
    """Infer application root from workbook path (data/seed/master_excel/*.xlsx)."""
    return workbook_path.resolve().parents[3]


def _validate_cross_sheet_links(
    customer_rows: list[tuple[int, dict[str, Any]]],
    user_rows: list[tuple[int, dict[str, Any]]],
    app_root: Path,
    issues: list[ValidationIssue],
) -> None:
    if not customer_rows:
        return

    customer_ids = {
        _cell_str(data.get("customer_id"))
        for _row_number, data in customer_rows
        if not _is_empty(data.get("customer_id"))
    }

    for row_number, data in user_rows:
        if _cell_str(data.get("role")) != "customer":
            continue
        customer_id = _cell_str(data.get("customer_id"))
        if customer_id and customer_id not in customer_ids:
            issues.append(
                ValidationIssue(
                    level="error",
                    sheet="users",
                    row=row_number,
                    column="customer_id",
                    code="MISSING_CUSTOMER_REF",
                    message=f"customer_id '{customer_id}' not found in customers sheet",
                )
            )

    for row_number, data in customer_rows:
        folder_path = _cell_str(data.get("folder_path"))
        if not folder_path:
            continue
        resolved = (app_root / folder_path).resolve()
        if not resolved.is_dir():
            issues.append(
                ValidationIssue(
                    level="error",
                    sheet="customers",
                    row=row_number,
                    column="folder_path",
                    code="MISSING_FOLDER_PATH",
                    message=f"folder_path does not exist: {folder_path}",
                )
            )


def _sheet_has_data(sheet_data_rows: dict[str, list[tuple[int, dict[str, Any]]]], sheet_name: str) -> bool:
    return bool(sheet_data_rows.get(sheet_name))


def _ids_from_sheet(
    sheet_data_rows: dict[str, list[tuple[int, dict[str, Any]]]],
    sheet_name: str,
    column: str,
) -> set[str]:
    values: set[str] = set()
    for _row_number, data in sheet_data_rows.get(sheet_name, []):
        raw = data.get(column)
        if not _is_empty(raw):
            values.add(_cell_str(raw))
    return values


_FK_ERROR_CODE_BY_PARENT_SHEET: dict[str, str] = {
    "customers": "MISSING_CUSTOMER_REF",
    "loan_applications": "MISSING_APPLICATION_REF",
    "loans": "MISSING_LOAN_REF",
    "payment_transactions": "MISSING_TRANSACTION_REF",
    "tickets": "MISSING_TICKET_REF",
    "users": "MISSING_USER_REF",
}

_LENDING_SHEETS = frozenset(
    {
        "loan_applications",
        "loans",
        "kyc_documents",
        "payment_transactions",
        "repayment_schedule",
    }
)

_BUREAU_OFFERS_SHEETS = frozenset(
    {
        "bureau_reporting_logs",
        "offers",
        "rm_mapping",
        "fraud_cases",
    }
)

_POLICY_SHEETS = frozenset({"knowledge_documents"})

_EVALUATION_SHEETS = frozenset({"evaluation_cases"})


def _validate_lending_foreign_keys(
    sheet_data_rows: dict[str, list[tuple[int, dict[str, Any]]]],
    issues: list[ValidationIssue],
) -> None:
    if not _sheet_has_data(sheet_data_rows, "loan_applications"):
        return

    parent_ids: dict[str, set[str]] = {}
    for parent_sheet in ("customers", "loan_applications", "loans", "payment_transactions"):
        if not _sheet_has_data(sheet_data_rows, parent_sheet):
            continue
        parent_spec = SHEET_SPECS_BY_NAME[parent_sheet]
        parent_ids[parent_sheet] = _ids_from_sheet(
            sheet_data_rows,
            parent_sheet,
            parent_spec.stable_id_column,
        )

    for sheet_name in _LENDING_SHEETS:
        if not _sheet_has_data(sheet_data_rows, sheet_name):
            continue
        spec = SHEET_SPECS_BY_NAME[sheet_name]
        for row_number, data in sheet_data_rows.get(sheet_name, []):
            for column, target in spec.foreign_keys.items():
                raw_value = data.get(column)
                if _is_empty(raw_value):
                    continue
                value = _cell_str(raw_value)
                parent_sheet, _parent_column = target.split(".", maxsplit=1)
                if parent_sheet not in parent_ids:
                    continue
                if value not in parent_ids[parent_sheet]:
                    issues.append(
                        ValidationIssue(
                            level="error",
                            sheet=sheet_name,
                            row=row_number,
                            column=column,
                            code=_FK_ERROR_CODE_BY_PARENT_SHEET.get(
                                parent_sheet,
                                "MISSING_FOREIGN_REF",
                            ),
                            message=(
                                f"{column} '{value}' not found in {parent_sheet} sheet"
                            ),
                        )
                    )


def _validate_bureau_offers_foreign_keys(
    sheet_data_rows: dict[str, list[tuple[int, dict[str, Any]]]],
    issues: list[ValidationIssue],
) -> None:
    if not any(_sheet_has_data(sheet_data_rows, sheet_name) for sheet_name in _BUREAU_OFFERS_SHEETS):
        return

    parent_ids: dict[str, set[str]] = {}
    for parent_sheet in ("customers", "loans", "payment_transactions", "tickets"):
        if not _sheet_has_data(sheet_data_rows, parent_sheet):
            continue
        parent_spec = SHEET_SPECS_BY_NAME[parent_sheet]
        parent_ids[parent_sheet] = _ids_from_sheet(
            sheet_data_rows,
            parent_sheet,
            parent_spec.stable_id_column,
        )

    for sheet_name in _BUREAU_OFFERS_SHEETS:
        if not _sheet_has_data(sheet_data_rows, sheet_name):
            continue
        spec = SHEET_SPECS_BY_NAME[sheet_name]
        for row_number, data in sheet_data_rows.get(sheet_name, []):
            for column, target in spec.foreign_keys.items():
                raw_value = data.get(column)
                if _is_empty(raw_value):
                    continue
                value = _cell_str(raw_value)
                parent_sheet, _parent_column = target.split(".", maxsplit=1)
                if parent_sheet == "tickets" and parent_sheet not in parent_ids:
                    continue
                if parent_sheet not in parent_ids:
                    continue
                if value not in parent_ids[parent_sheet]:
                    issues.append(
                        ValidationIssue(
                            level="error",
                            sheet=sheet_name,
                            row=row_number,
                            column=column,
                            code=_FK_ERROR_CODE_BY_PARENT_SHEET.get(
                                parent_sheet,
                                "MISSING_FOREIGN_REF",
                            ),
                            message=(
                                f"{column} '{value}' not found in {parent_sheet} sheet"
                            ),
                        )
                    )


def _validate_policy_foreign_keys(
    sheet_data_rows: dict[str, list[tuple[int, dict[str, Any]]]],
    issues: list[ValidationIssue],
) -> None:
    if not _sheet_has_data(sheet_data_rows, "knowledge_documents"):
        return

    parent_ids: dict[str, set[str]] = {}
    if _sheet_has_data(sheet_data_rows, "users"):
        parent_spec = SHEET_SPECS_BY_NAME["users"]
        parent_ids["users"] = _ids_from_sheet(
            sheet_data_rows,
            "users",
            parent_spec.stable_id_column,
        )

    spec = SHEET_SPECS_BY_NAME["knowledge_documents"]
    for row_number, data in sheet_data_rows.get("knowledge_documents", []):
        for column, target in spec.foreign_keys.items():
            raw_value = data.get(column)
            if _is_empty(raw_value):
                continue
            value = _cell_str(raw_value)
            parent_sheet, _parent_column = target.split(".", maxsplit=1)
            if parent_sheet not in parent_ids:
                continue
            if value not in parent_ids[parent_sheet]:
                issues.append(
                    ValidationIssue(
                        level="error",
                        sheet="knowledge_documents",
                        row=row_number,
                        column=column,
                        code=_FK_ERROR_CODE_BY_PARENT_SHEET.get(
                            parent_sheet,
                            "MISSING_FOREIGN_REF",
                        ),
                        message=(
                            f"{column} '{value}' not found in {parent_sheet} sheet"
                        ),
                    )
                )


def _validate_policy_storage_paths(
    sheet_data_rows: dict[str, list[tuple[int, dict[str, Any]]]],
    app_root: Path,
    issues: list[ValidationIssue],
) -> None:
    if not _sheet_has_data(sheet_data_rows, "knowledge_documents"):
        return

    for row_number, data in sheet_data_rows.get("knowledge_documents", []):
        storage_path = _cell_str(data.get("storage_path"))
        if not storage_path:
            continue
        resolved = (app_root / storage_path).resolve()
        if not resolved.is_file():
            issues.append(
                ValidationIssue(
                    level="error",
                    sheet="knowledge_documents",
                    row=row_number,
                    column="storage_path",
                    code="MISSING_POLICY_FILE",
                    message=f"storage_path does not exist: {storage_path}",
                )
            )


def _validate_evaluation_foreign_keys(
    sheet_data_rows: dict[str, list[tuple[int, dict[str, Any]]]],
    issues: list[ValidationIssue],
) -> None:
    if not _sheet_has_data(sheet_data_rows, "evaluation_cases"):
        return

    parent_ids: dict[str, set[str]] = {}
    if _sheet_has_data(sheet_data_rows, "customers"):
        parent_spec = SHEET_SPECS_BY_NAME["customers"]
        parent_ids["customers"] = _ids_from_sheet(
            sheet_data_rows,
            "customers",
            parent_spec.stable_id_column,
        )

    spec = SHEET_SPECS_BY_NAME["evaluation_cases"]
    for row_number, data in sheet_data_rows.get("evaluation_cases", []):
        for column, target in spec.foreign_keys.items():
            raw_value = data.get(column)
            if _is_empty(raw_value):
                continue
            value = _cell_str(raw_value)
            parent_sheet, _parent_column = target.split(".", maxsplit=1)
            if parent_sheet not in parent_ids:
                continue
            if value not in parent_ids[parent_sheet]:
                issues.append(
                    ValidationIssue(
                        level="error",
                        sheet="evaluation_cases",
                        row=row_number,
                        column=column,
                        code=_FK_ERROR_CODE_BY_PARENT_SHEET.get(
                            parent_sheet,
                            "MISSING_FOREIGN_REF",
                        ),
                        message=(
                            f"{column} '{value}' not found in {parent_sheet} sheet"
                        ),
                    )
                )


def validate_workbook(workbook_path: Path) -> ValidationResult:
    """Validate seed workbook structure and row-level constraints when data exists."""
    issues: list[ValidationIssue] = []
    sheets_checked = 0
    rows_checked = 0

    if not workbook_path.is_file():
        issues.append(
            ValidationIssue(
                level="error",
                sheet=None,
                row=None,
                column=None,
                code="WORKBOOK_NOT_FOUND",
                message=f"workbook not found: {workbook_path}",
            )
        )
        return ValidationResult(
            ok=False,
            issues=issues,
            workbook_path=workbook_path,
        )

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        issues.append(
            ValidationIssue(
                level="error",
                sheet=None,
                row=None,
                column=None,
                code="OPENPYXL_MISSING",
                message="openpyxl is required; run uv sync",
            )
        )
        return ValidationResult(
            ok=False,
            issues=issues,
            workbook_path=workbook_path,
        )

    required_sheets = sheets_required_for_validate()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    customer_data_rows: list[tuple[int, dict[str, Any]]] = []
    user_data_rows: list[tuple[int, dict[str, Any]]] = []
    sheet_data_rows: dict[str, list[tuple[int, dict[str, Any]]]] = {}
    app_root = _resolve_app_root(workbook_path)
    try:
        present_sheets = set(workbook.sheetnames)
        for sheet_name in required_sheets:
            if sheet_name not in present_sheets:
                issues.append(
                    ValidationIssue(
                        level="error",
                        sheet=sheet_name,
                        row=None,
                        column=None,
                        code="MISSING_SHEET",
                        message=f"required sheet '{sheet_name}' is missing",
                    )
                )

        extra_sheets = present_sheets - set(required_sheets)
        for sheet_name in sorted(extra_sheets):
            issues.append(
                ValidationIssue(
                    level="warning",
                    sheet=sheet_name,
                    row=None,
                    column=None,
                    code="EXTRA_SHEET",
                    message=f"unexpected extra sheet '{sheet_name}'",
                )
            )

        for sheet_name in required_sheets:
            if sheet_name not in present_sheets:
                continue

            spec = SHEET_SPECS_BY_NAME[sheet_name]
            worksheet = workbook[sheet_name]
            rows = list(worksheet.iter_rows(values_only=True))
            sheets_checked += 1

            if not rows:
                issues.append(
                    ValidationIssue(
                        level="error",
                        sheet=sheet_name,
                        row=None,
                        column=None,
                        code="EMPTY_SHEET",
                        message="sheet has no header row",
                    )
                )
                continue

            headers = _parse_header_row(rows[0])
            if not headers:
                issues.append(
                    ValidationIssue(
                        level="error",
                        sheet=sheet_name,
                        row=1,
                        column=None,
                        code="MISSING_COLUMN",
                        message="header row is empty",
                    )
                )
                continue

            _validate_headers(spec, headers, issues)

            if len(rows) <= 1:
                continue

            id_values: list[tuple[int, str]] = []
            for row_index, row in enumerate(rows[1:], start=2):
                if all(_is_empty(cell) for cell in row):
                    continue
                rows_checked += 1
                row_data = _row_dict(headers, row)
                _validate_row(spec, headers, row_index, row, issues)
                stable_raw = row_data.get(spec.stable_id_column)
                if not _is_empty(stable_raw):
                    id_values.append((row_index, _cell_str(stable_raw)))
                if sheet_name == "customers":
                    customer_data_rows.append((row_index, row_data))
                elif sheet_name == "users":
                    user_data_rows.append((row_index, row_data))
                sheet_data_rows.setdefault(sheet_name, []).append((row_index, row_data))

            _validate_duplicate_ids(spec, id_values, issues)

        _validate_cross_sheet_links(
            customer_data_rows,
            user_data_rows,
            app_root,
            issues,
        )
        _validate_lending_foreign_keys(sheet_data_rows, issues)
        _validate_bureau_offers_foreign_keys(sheet_data_rows, issues)
        _validate_policy_foreign_keys(sheet_data_rows, issues)
        _validate_policy_storage_paths(sheet_data_rows, app_root, issues)
        _validate_evaluation_foreign_keys(sheet_data_rows, issues)
    finally:
        workbook.close()

    has_errors = any(issue.level == "error" for issue in issues)
    return ValidationResult(
        ok=not has_errors,
        issues=issues,
        workbook_path=workbook_path,
        sheets_checked=sheets_checked,
        rows_checked=rows_checked,
    )
